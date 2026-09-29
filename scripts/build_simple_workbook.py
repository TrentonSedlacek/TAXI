"""Build analysis/Theta Xi GPA and Recruitment.xlsx, the version to share with the chapter.

Run from the repo root:  python3 scripts/build_simple_workbook.py
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "analysis" / "Theta Xi GPA and Recruitment.xlsx"
US = "Theta Xi"
FIRST = "Fall 2018"      # first semester with every IFC house
FIRST_YEAR = 2011        # first full exec year in our records
SPARE = 4

# colors (NE DHHS palette)
BLUE, GOLD, GRAY, LIME, MIST, BERRY = "00607F", "FFC843", "4D4D4F", "BABF33", "B9C8D3", "BB1F53"

RUSH_CHAIRS = {2015: "Tyler Brattain", 2016: "Tyler Brattain", 2017: "Nick Mason", 2018: "Brennan, Evan Cox",
               2019: "Eli Soell, Austin Bennet", 2020: "Tyler Lechner, Eli Soell",
               2021: "Max Dikker, Connor Wieseman", 2022: "Elijah Doyle, Michael Leiting",
               2023: "Ben Morse, Owen Micek", 2024: "Sebastian, Miles"}

F = Font(name="Arial", size=10)
FI = Font(name="Arial", size=10, italic=True, color=GRAY)
B = Font(name="Arial", size=10, bold=True)
T = Font(name="Arial", size=16, bold=True, color=BLUE)
SUB = Font(name="Arial", size=9, color=GRAY)
HF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor=BLUE)
USFILL = PatternFill("solid", fgColor=GOLD)
LINE = Side(style="thin", color=MIST)
BOX = Border(top=LINE, bottom=LINE, left=LINE, right=LINE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right")


def key(t):
    s, y = t.split()
    return int(y) * 2 + (s == "Fall")


def num(v):
    if v in (None, ""):
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


with open(ROOT / "data" / "chapter_semester_data.csv", newline="") as fh:
    rows = [r for r in csv.DictReader(fh) if r["council"] == "IFC" and
            (key(r["semester"]) >= key(FIRST) or r["chapter"] == US)]
with open(ROOT / "data" / "campus_stats.csv", newline="") as fh:
    campus = list(csv.DictReader(fh))
rows.sort(key=lambda r: (key(r["semester"]), r["chapter"]))
peer_terms = sorted({r["semester"] for r in rows if key(r["semester"]) >= key(FIRST)}, key=key)
latest = peer_terms[-1]
last_fall_year = max(int(t.split()[1]) for t in peer_terms if t.startswith("Fall"))

wb = Workbook()
sem_ws = wb.active
sem_ws.title = "By Semester"
rec_ws = wb.create_sheet("Recruitment by Year")
gpa_ws = wb.create_sheet("GPA by Year")
house_ws = wb.create_sheet("IFC Houses")
data_ws = wb.create_sheet("Data")
avg_ws = wb.create_sheet("UNL Averages")


def cell(ws, r, c, v, font=F, fmt=None, fill=None, border=True):
    x = ws.cell(r, c, v)
    x.font = font
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    if border:
        x.border = BOX
    return x


def header(ws, r, labels, height=30):
    for i, text in enumerate(labels, start=1):
        cell(ws, r, i, text, HF, fill=HFILL).alignment = CENTER
    ws.row_dimensions[r].height = height


def title(ws, text, sub):
    cell(ws, 1, 1, text, T, border=False)
    cell(ws, 2, 1, sub, SUB, border=False)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


def best_worst(ws, rng, higher_is_better=True):
    good = DifferentialStyle(fill=PatternFill("solid", fgColor=LIME, bgColor=LIME), font=Font(bold=True))
    bad = DifferentialStyle(fill=PatternFill("solid", fgColor=BERRY, bgColor=BERRY),
                            font=Font(bold=True, color="FFFFFF"))
    ws.conditional_formatting.add(rng, Rule(type="top10", rank=1, bottom=not higher_is_better, dxf=good))
    ws.conditional_formatting.add(rng, Rule(type="top10", rank=1, bottom=higher_is_better, dxf=bad))


def color_series(chart, colors):
    for s, c in zip(chart.series, colors):
        s.graphicalProperties.solidFill = c
        s.graphicalProperties.line.solidFill = c


def neg_red(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                  font=Font(name="Arial", color=BERRY)))


# ------------------------------------------------------------------ Data
header(data_ws, 1, ["Semester", "House", "Members", "New Members", "GPA", "Source"], height=20)
for i, r in enumerate(rows, start=2):
    ours_only = r.get("from_records", "")
    src = "UNL" if not ours_only else ("Our records" if ours_only == "all" else "UNL + our records")
    font = FI if ours_only == "all" else F
    cell(data_ws, i, 1, r["semester"], font)
    cell(data_ws, i, 2, r["chapter"], font)
    cell(data_ws, i, 3, num(r["total_members"]), font, "0")
    cell(data_ws, i, 4, num(r["new_members"]), FI if "new_members" in ours_only else font, "0")
    cell(data_ws, i, 5, num(r["chapter_gpa"]), FI if "chapter_gpa" in ours_only else font, "0.000")
    cell(data_ws, i, 6, src, font)
data_ws.freeze_panes = "A2"
data_ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
widths(data_ws, [13, 22, 10, 12, 8, 17])
data_notes = ["IFC houses from UNL's report cards, Fall 2018 on. Theta Xi before that is from our own records.",
              "Gray italics = our own records. UNL left these blank for us: Spring 2020 GPA and new members,",
              "Spring 2022 new members, Fall 2023 new members.",
              "Fall 2023 has no UNL scorecard, so that semester is from UNL's grade report.",
              "To add a semester, paste the new IFC rows at the bottom."]
for k, text in enumerate(data_notes):
    cell(data_ws, 1 + k, 8, text, SUB, border=False)

N = 1500
SEM, HOUSE, MEM, NM, GPA = (f"Data!${c}$2:${c}${N}" for c in "ABCDE")

# ------------------------------------------------------------------ UNL Averages
header(avg_ws, 1, ["Semester", "All UNL Men", "All Greek Men", "All UNL Students"], height=20)
for i, r in enumerate(campus, start=2):
    cell(avg_ws, i, 1, r["semester"])
    cell(avg_ws, i, 2, num(r["All Male"]), fmt="0.000")
    cell(avg_ws, i, 3, num(r["Greek Male"]), fmt="0.000")
    cell(avg_ws, i, 4, num(r["All UNL"]), fmt="0.000")
widths(avg_ws, [13, 12, 13, 15])
cell(avg_ws, 1, 6, "From UNL's All-Community Grade Report each semester.", SUB, border=False)
avg_ws.freeze_panes = "A2"


def ours(rng, sem):
    return (f'IF(COUNTIFS({SEM},{sem},{HOUSE},"{US}",{rng},"<>")=0,"",'
            f'SUMIFS({rng},{SEM},{sem},{HOUSE},"{US}"))')


def ifc_avg(rng, sem):
    return f'IF(COUNTIFS({SEM},{sem},{rng},"<>")<5,"",AVERAGEIFS({rng},{SEM},{sem}))'


def unl(col, sem):
    look = f"INDEX('UNL Averages'!${col}:${col},MATCH({sem},'UNL Averages'!$A:$A,0))"
    return f'IFERROR(IF({look}="","",{look}),"")'


# ------------------------------------------------------------------ By Semester
title(sem_ws, "Theta Xi vs the IFC", "Every semester since Fall 2018. IFC Avg = the average IFC house that semester.")
header(sem_ws, 4, ["Semester", "TX GPA", "Greek Men GPA", "TX vs Greek Men", "All UNL Men GPA", "TX Rank in IFC",
                   "TX Members", "IFC Avg Members", "TX New Members", "IFC Avg New Members", "TX Rank in IFC"],
       height=42)
sem_ws.freeze_panes = "B5"
for i in range(len(peer_terms) + SPARE):
    r = 5 + i
    a = f"$A{r}"
    cell(sem_ws, r, 1, peer_terms[i] if i < len(peer_terms) else None, B)
    cell(sem_ws, r, 2, f'=IF({a}="","",{ours(GPA, a)})', fmt="0.000")
    cell(sem_ws, r, 3, f'=IF({a}="","",{unl("C", a)})', fmt="0.000")
    cell(sem_ws, r, 4, f'=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})', fmt="+0.000;-0.000;0.000")
    cell(sem_ws, r, 5, f'=IF({a}="","",{unl("B", a)})', fmt="0.000")
    cell(sem_ws, r, 6, f'=IF(OR(B{r}="",COUNTIFS({SEM},{a},{GPA},">0")<5),"",'
                       f'COUNTIFS({SEM},{a},{GPA},">"&B{r})+1&" of "&COUNTIFS({SEM},{a},{GPA},">0"))').alignment = RIGHT
    cell(sem_ws, r, 7, f'=IF({a}="","",{ours(MEM, a)})', fmt="0")
    cell(sem_ws, r, 8, f'=IF({a}="","",{ifc_avg(MEM, a)})', fmt="0")
    cell(sem_ws, r, 9, f'=IF({a}="","",{ours(NM, a)})', fmt="0")
    cell(sem_ws, r, 10, f'=IF({a}="","",{ifc_avg(NM, a)})', fmt="0.0")
    cell(sem_ws, r, 11, f'=IF(OR(I{r}="",J{r}=""),"",'
                        f'COUNTIFS({SEM},{a},{NM},">"&I{r})+1&" of "&COUNTIFS({SEM},{a},{NM},">=0"))').alignment = RIGHT
last = 4 + len(peer_terms) + SPARE
neg_red(sem_ws, f"D5:D{last}")
widths(sem_ws, [13, 8, 9, 9, 9, 9, 9, 9, 9, 9, 9])
cell(sem_ws, last + 1, 1, "Blank = not reported. Spring 2020 was pass/no-pass, so UNL has no GPAs for other houses.",
     SUB, border=False)
known = 4 + len(peer_terms)
cats = Reference(sem_ws, min_col=1, min_row=5, max_row=known)
c1 = BarChart()
c1.title, c1.height, c1.width = "GPA", 7.5, 17
c1.y_axis.scaling.min, c1.y_axis.scaling.max = 2.6, 3.6
c1.y_axis.number_format = "0.0"
for col in (2, 3):
    c1.add_data(Reference(sem_ws, min_col=col, min_row=4, max_row=known), titles_from_data=True)
c1.set_categories(cats)
color_series(c1, [BLUE, GOLD])
sem_ws.add_chart(c1, f"A{last + 3}")
c2 = BarChart()
c2.title, c2.height, c2.width = "New Members", 7.5, 17
for col in (9, 10):
    c2.add_data(Reference(sem_ws, min_col=col, min_row=4, max_row=known), titles_from_data=True)
c2.set_categories(cats)
color_series(c2, [BLUE, GOLD])
sem_ws.add_chart(c2, f"G{last + 3}")

# ------------------------------------------------------------------ Recruitment by Year
years = list(range(FIRST_YEAR, last_fall_year + 2))          # through the current (partial) year
title(rec_ws, "Recruitment by Year", "Each year = spring + fall recruitment, one exec board "
                                     "(elected in November). Green = best, red = worst.")
header(rec_ws, 4, ["Year", "Rush Chairs", "New Members", "IFC Avg New Members", "TX vs Avg House",
                   "Members Last Fall", "Class Size vs House", "Members This Fall", "Change"], height=42)
rec_ws.freeze_panes = "C5"
for i, y in enumerate(years):
    r = 5 + i
    sp, fa, pf = f'"Spring "&$A{r}', f'"Fall "&$A{r}', f'"Fall "&($A{r}-1)'
    cell(rec_ws, r, 1, y, B)
    cell(rec_ws, r, 2, RUSH_CHAIRS.get(y))
    cell(rec_ws, r, 3, f'=IF(OR({ours(NM, sp)}="",{ours(NM, fa)}=""),"",{ours(NM, sp)}+{ours(NM, fa)})', fmt="0")
    cell(rec_ws, r, 4, f'=IF(OR({ifc_avg(NM, sp)}="",{ifc_avg(NM, fa)}=""),"",{ifc_avg(NM, sp)}+{ifc_avg(NM, fa)})',
         fmt="0.0")
    cell(rec_ws, r, 5, f'=IF(OR(C{r}="",D{r}=""),"",C{r}/D{r})', fmt="0%")
    cell(rec_ws, r, 6, f"={ours(MEM, pf)}", fmt="0")
    cell(rec_ws, r, 7, f'=IF(OR(C{r}="",F{r}=""),"",C{r}/F{r})', fmt="0%")
    cell(rec_ws, r, 8, f"={ours(MEM, fa)}", fmt="0")
    cell(rec_ws, r, 9, f'=IF(OR(F{r}="",H{r}=""),"",H{r}-F{r})', fmt="+0;-0;0")
done = 4 + (last_fall_year - FIRST_YEAR + 1)
for col in "EGI":
    best_worst(rec_ws, f"{col}5:{col}{done}")
neg_red(rec_ws, f"I5:I{done}")
widths(rec_ws, [8, 26, 10, 11, 10, 10, 10, 10, 9])
rec_notes = ["TX vs Avg House: 100% = as many new members as the average IFC house. Only from 2019 (needs every "
             "house's numbers).",
             "Class Size vs House: new members that year divided by the house size the fall before.",
             "Rush chair names are from the AE List sheet. Add missing ones in column B."]
for k, text in enumerate(rec_notes):
    cell(rec_ws, 5 + len(years) + k, 1, text, SUB, border=False)
ycats = Reference(rec_ws, min_col=1, min_row=5, max_row=done)
c3 = BarChart()
c3.title, c3.height, c3.width = "New Members per Year", 7.5, 16
c3.add_data(Reference(rec_ws, min_col=3, min_row=4, max_row=done), titles_from_data=True)
c3.set_categories(ycats)
color_series(c3, [BLUE])
c3.legend = None
rec_ws.add_chart(c3, f"A{9 + len(years)}")
c4 = BarChart()
c4.title, c4.height, c4.width = "Members Each Fall", 7.5, 16
c4.add_data(Reference(rec_ws, min_col=8, min_row=4, max_row=done), titles_from_data=True)
c4.set_categories(ycats)
color_series(c4, [GOLD])
c4.legend = None
rec_ws.add_chart(c4, f"E{9 + len(years)}")

# ------------------------------------------------------------------ GPA by Year
title(gpa_ws, "GPA by Year", "Each year = spring + fall semesters, one exec board. Compared with all Greek men "
                             "that year, since grades across campus rise and fall. Green = best, red = worst.")
header(gpa_ws, 4, ["Year", "Scholarship Chair", "Spring GPA", "Fall GPA", "TX Year GPA", "Greek Men GPA",
                   "TX vs Greek Men", "Spring Rank in IFC", "Fall Rank in IFC"], height=42)
gpa_ws.freeze_panes = "C5"
for i, y in enumerate(years):
    r = 5 + i
    sp, fa = f'"Spring "&$A{r}', f'"Fall "&$A{r}'
    cell(gpa_ws, r, 1, y, B)
    cell(gpa_ws, r, 2, None)
    cell(gpa_ws, r, 3, f"={ours(GPA, sp)}", fmt="0.000")
    cell(gpa_ws, r, 4, f"={ours(GPA, fa)}", fmt="0.000")
    cell(gpa_ws, r, 5, f'=IF(COUNT(C{r}:D{r})<2,"",AVERAGE(C{r}:D{r}))', fmt="0.000")
    cell(gpa_ws, r, 6, f'=IF(OR({unl("C", sp)}="",{unl("C", fa)}=""),"",({unl("C", sp)}+{unl("C", fa)})/2)',
         fmt="0.000")
    cell(gpa_ws, r, 7, f'=IF(OR(E{r}="",F{r}=""),"",E{r}-F{r})', fmt="+0.000;-0.000;0.000")
    for col, sem, g in ((8, sp, "C"), (9, fa, "D")):
        cell(gpa_ws, r, col, f'=IF(OR({g}{r}="",COUNTIFS({SEM},{sem},{GPA},">0")<5),"",'
                             f'COUNTIFS({SEM},{sem},{GPA},">"&{g}{r})+1&" of "&COUNTIFS({SEM},{sem},{GPA},">0"))'
             ).alignment = RIGHT
best_worst(gpa_ws, f"G5:G{done}")
neg_red(gpa_ws, f"G5:G{done}")
widths(gpa_ws, [8, 20, 9, 9, 10, 10, 10, 10, 10])
cell(gpa_ws, 5 + len(years), 1, "Add scholarship chairs in column B. IFC ranks start in Fall 2018.", SUB, border=False)
c5 = BarChart()
c5.title, c5.height, c5.width = "TX GPA vs Greek Men, by Year", 7.5, 16
c5.add_data(Reference(gpa_ws, min_col=7, min_row=4, max_row=done), titles_from_data=True)
c5.set_categories(Reference(gpa_ws, min_col=1, min_row=5, max_row=done))
color_series(c5, [BLUE])
c5.legend = None
c5.x_axis.tickLblPos = "low"
c5.y_axis.number_format = "+0.00;-0.00;0.00"
gpa_ws.add_chart(c5, f"A{8 + len(years)}")

# ------------------------------------------------------------------ IFC Houses
title(house_ws, "Every IFC House", "Change the two blue cells to see a different semester or year.")
cell(house_ws, 3, 1, "GPA and members for", B, border=False)
cell(house_ws, 3, 3, latest, Font(name="Arial", size=10, bold=True, color="0000FF"))
cell(house_ws, 4, 1, "New members for year", B, border=False)
cell(house_ws, 4, 3, last_fall_year, Font(name="Arial", size=10, bold=True, color="0000FF"))
header(house_ws, 6, ["House", "", "GPA", "GPA Rank", "Members", "New Members (Spring + Fall)", "New Member Rank"],
       height=42)
house_ws.merge_cells("A6:B6")
current = [r for r in rows if r["semester"] == latest]
current.sort(key=lambda r: -(num(r["chapter_gpa"]) or 0))
h0, h1 = 7, 7 + len(current) - 1
sp, fa = '"Spring "&$C$4', '"Fall "&$C$4'
for i, r in enumerate(current):
    rr = h0 + i
    us = r["chapter"] == US
    fill, font = (USFILL, B) if us else (None, F)
    house_ws.merge_cells(f"A{rr}:B{rr}")
    cell(house_ws, rr, 1, r["chapter"], font, fill=fill)
    cell(house_ws, rr, 2, None, font, fill=fill)
    a = f"$A{rr}"
    cell(house_ws, rr, 3, f'=IF(COUNTIFS({SEM},$C$3,{HOUSE},{a},{GPA},"<>")=0,"",SUMIFS({GPA},{SEM},$C$3,{HOUSE},{a}))',
         font, "0.000", fill)
    cell(house_ws, rr, 4, f'=IF(C{rr}="","",COUNTIF($C${h0}:$C${h1},">"&C{rr})+1)', font, "0", fill)
    cell(house_ws, rr, 5, f'=IF(COUNTIFS({SEM},$C$3,{HOUSE},{a},{MEM},"<>")=0,"",SUMIFS({MEM},{SEM},$C$3,{HOUSE},{a}))',
         font, "0", fill)
    cell(house_ws, rr, 6, f'=IF(OR(COUNTIFS({SEM},{sp},{HOUSE},{a},{NM},"<>")=0,COUNTIFS({SEM},{fa},{HOUSE},{a},{NM},"<>")=0),"",'
                          f'SUMIFS({NM},{SEM},{sp},{HOUSE},{a})+SUMIFS({NM},{SEM},{fa},{HOUSE},{a}))', font, "0", fill)
    cell(house_ws, rr, 7, f'=IF(F{rr}="","",COUNTIF($F${h0}:$F${h1},">"&F{rr})+1)', font, "0", fill)
widths(house_ws, [14, 12, 9, 9, 10, 14, 11])
cell(house_ws, h1 + 2, 1, f"Houses on the {latest} report, sorted by {latest} GPA.", SUB, border=False)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("Data", "UNL Averages")
    for ch in ws._charts:
        if ch.legend is not None:
            ch.legend.position = "b"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.tabColor = BLUE if ws.title not in ("Data", "UNL Averages") else MIST
wb.save(OUT)
print("saved", OUT)
