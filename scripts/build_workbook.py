"""Build analysis/Theta Xi vs UNL Fraternities.xlsx from the tidy CSVs in data/.

Run from the repo root:  python3 scripts/build_workbook.py
The workbook is formula driven, so after it is built you can keep it current
in Excel by pasting new rows into the Chapter Data and Campus Stats tabs.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "analysis" / "Theta Xi vs UNL Fraternities.xlsx"

OUR_CHAPTER = "Theta Xi"
FIRST_PEER_TERM = "Fall 2018"   # first semester with every chapter's numbers
EXTRA_TERMS = 4                 # empty future columns/rows ready for new data
DATA_ROWS = 2500                # Chapter Data rows the formulas read (and Term # is pre-filled)

# ---------------------------------------------------------------- styles
FONT = "Arial"
F = Font(name=FONT, size=10)
F_B = Font(name=FONT, size=10, bold=True)
F_IN = Font(name=FONT, size=10, color="0000FF")          # typed-in data
F_LINK = Font(name=FONT, size=10, color="008000")        # pulls from another tab
F_H = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_T = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_US = PatternFill("solid", fgColor="FFF2CC")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_GAP = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

GPA = "0.000"
INT = "0"
PCT = "0%"
IDX = "0"
SIGNED = "+0.000;-0.000;0.000"
SIGNED_PCT = "+0%;-0%;0%"


def head(ws, row, labels, col=1, height=None):
    for i, text in enumerate(labels):
        c = ws.cell(row, col + i, text)
        c.font, c.fill, c.alignment, c.border = F_H, FILL_H, CENTER, BOX
    if height:
        ws.row_dimensions[row].height = height


def put(ws, row, col, value, font=F, fmt=None, fill=None):
    c = ws.cell(row, col, value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def read_csv(name):
    with open(DATA / name, newline="") as fh:
        return list(csv.DictReader(fh))


def num(v):
    if v is None or str(v).strip() == "":
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def term_key(term):
    season, year = term.split()
    return int(year) * 2 + (1 if season == "Fall" else 0)


def next_term(term):
    season, year = term.split()
    return f"Spring {int(year) + 1}" if season == "Fall" else f"Fall {year}"



from workbook_text import EVALUATION, OLD_INDEX_NOTES, README  # noqa: E402  (long text lives there)

# ---------------------------------------------------------------- data
chapter_rows = read_csv("chapter_semester_data.csv")
campus_rows = read_csv("campus_stats.csv")

known_terms = sorted({r["semester"] for r in chapter_rows} | {r["semester"] for r in campus_rows},
                     key=term_key)
all_terms = [known_terms[0]]
while all_terms[-1] != known_terms[-1]:
    all_terms.append(next_term(all_terms[-1]))
for _ in range(EXTRA_TERMS):
    all_terms.append(next_term(all_terms[-1]))
peer_terms = [t for t in all_terms if term_key(t) >= term_key(FIRST_PEER_TERM)]
latest_term = known_terms[-1]

ifc_chapters = sorted({r["chapter"] for r in chapter_rows
                       if r["council"] == "IFC" and term_key(r["semester"]) >= term_key(FIRST_PEER_TERM)})

wb = Workbook()
wb.remove(wb.active)
tabs = {}
for name in ["Read Me", "Summary", "TX vs IFC", "Combined Score", "Recruitment Index",
             "GPA Matrix", "Members Matrix", "New Members Matrix",
             "TX History", "Chapter Data", "Campus Stats", "IFC Size (legacy)", "Lists", "Old Index Review"]:
    tabs[name] = wb.create_sheet(name)

CD = "'Chapter Data'!"

def cd(col):
    return f"{CD}${col}$2:${col}${DATA_ROWS}"


CD_TERM, CD_ORDER, CD_COUNCIL, CD_CH = cd("A"), cd("B"), cd("C"), cd("D")
CD_MEM, CD_NM, CD_GPA = cd("E"), cd("F"), cd("G")
CD_RA, CD_RN = cd("J"), cd("K")

# ---------------------------------------------------------------- Lists
ws = tabs["Lists"]
head(ws, 1, ["Semester (in order)", "Term #", "", "IFC chapters", "", "Setting", "Value"])
for i, t in enumerate(all_terms, start=2):
    put(ws, i, 1, t, F_IN)
    put(ws, i, 2, i - 1)
for i, ch in enumerate(ifc_chapters, start=2):
    put(ws, i, 4, ch, F_IN)
put(ws, 2, 6, "Our chapter", F_B)
put(ws, 2, 7, OUR_CHAPTER, F_IN, fill=FILL_INPUT)
put(ws, 3, 6, "Latest semester to report", F_B)
put(ws, 3, 7, latest_term, F_IN, fill=FILL_INPUT)
put(ws, 4, 6, "Min. chapters before IFC stats show", F_B)
put(ws, 4, 7, 5, F_IN, fill=FILL_INPUT)
put(ws, 6, 6, "Add future semesters to column A and new chapters to column D. "
              "Matrices read the chapter list from column D.", F_NOTE)
ws.column_dimensions["A"].width = 20
ws.column_dimensions["D"].width = 24
ws.column_dimensions["F"].width = 36
ws.column_dimensions["G"].width = 16
wb.defined_names["OurChapter"] = DefinedName("OurChapter", attr_text="Lists!$G$2")
wb.defined_names["LatestTerm"] = DefinedName("LatestTerm", attr_text="Lists!$G$3")
wb.defined_names["MinChapters"] = DefinedName("MinChapters", attr_text="Lists!$G$4")
TERM_LIST = f"Lists!$A$2:$A${len(all_terms) + 1}"

# ---------------------------------------------------------------- Chapter Data
ws = tabs["Chapter Data"]
cols = ["Semester", "Term #", "Council", "Chapter", "Total Members", "New Members",
        "Chapter GPA", "Community GPA Rank", "Council GPA Rank", "Retention of Actives",
        "Retention of New Members", "Source", "Notes"]
head(ws, 1, cols, height=30)
ws.freeze_panes = "A2"
chapter_rows.sort(key=lambda r: (term_key(r["semester"]), r["council"], r["chapter"]))
for i, r in enumerate(chapter_rows, start=2):
    put(ws, i, 1, r["semester"], F_IN)
    put(ws, i, 3, r["council"], F_IN)
    put(ws, i, 4, r["chapter"], F_IN)
    put(ws, i, 5, num(r["total_members"]), F_IN, INT)
    put(ws, i, 6, num(r["new_members"]), F_IN, INT)
    put(ws, i, 7, num(r["chapter_gpa"]), F_IN, GPA)
    put(ws, i, 8, num(r["community_rank"]), F_IN, INT)
    put(ws, i, 9, num(r["council_rank"]), F_IN, INT)
    put(ws, i, 10, num(r["retention_active"]), F_IN, PCT)
    put(ws, i, 11, num(r["retention_new"]), F_IN, PCT)
    put(ws, i, 12, r["source_file"], F_IN)
    put(ws, i, 13, r.get("notes", ""), F_IN)
for i in range(2, DATA_ROWS + 1):
    put(ws, i, 2, f'=IF(A{i}="","",IFERROR(MATCH(A{i},{TERM_LIST},0),""))', F)
for col, w in zip("ABCDEFGHIJKLM", [13, 7, 8, 26, 9, 9, 9, 10, 9, 10, 11, 44, 50]):
    ws.column_dimensions[col].width = w
ws.auto_filter.ref = f"A1:M{len(chapter_rows) + 1}"

# ---------------------------------------------------------------- Campus Stats
ws = tabs["Campus Stats"]
campus_cols = ["All Male", "All Female", "All UNL", "Greek Male", "Greek Female", "All Greek",
               "Non-Greek Male", "Non-Greek Female", "All Non-Greek"]
head(ws, 1, ["Semester", "Term #"] + campus_cols + ["Source / notes"], height=30)
ws.freeze_panes = "C2"
by_term = {r["semester"]: r for r in campus_rows}
for i, t in enumerate(all_terms, start=2):
    put(ws, i, 1, t, F_IN)
    put(ws, i, 2, f"=MATCH(A{i},{TERM_LIST},0)")
    r = by_term.get(t, {})
    for j, c in enumerate(campus_cols, start=3):
        v = num(r.get(c))
        cell = put(ws, i, j, v, F_IN, GPA)
        if v is None and term_key(t) <= term_key(latest_term):
            cell.fill = FILL_GAP
    put(ws, i, 3 + len(campus_cols), r.get("notes", ""), F_IN)
ws.column_dimensions["A"].width = 13
ws.column_dimensions["B"].width = 7
for j in range(3, 3 + len(campus_cols)):
    ws.column_dimensions[L(j)].width = 11
ws.column_dimensions[L(3 + len(campus_cols))].width = 70
note_row = len(all_terms) + 3
put(ws, note_row, 1, "Orange cells: semesters where the UNL all-campus numbers are not in this repo. "
                     "They come from the FSL All-Community Academic Report (fsl.unl.edu); type them in "
                     "and every chart and gap column updates.", F_NOTE)
CS_ROWS = f"$A$2:$A${len(all_terms) + 1}"


def campus(term_ref, col_name):
    col = L(3 + campus_cols.index(col_name))
    return (f"IFERROR(IF(INDEX('Campus Stats'!${col}$2:${col}${len(all_terms) + 1},"
            f"MATCH({term_ref},'Campus Stats'!{CS_ROWS},0))=\"\",\"\","
            f"INDEX('Campus Stats'!${col}$2:${col}${len(all_terms) + 1},"
            f"MATCH({term_ref},'Campus Stats'!{CS_ROWS},0))),\"\")")


# ---------------------------------------------------------------- matrices
n_ch = len(ifc_chapters) + 5             # five spare rows for new chapters
first, last = 3, 3 + n_ch - 1            # chapter rows
mcols = len(peer_terms)
lastc = L(1 + mcols)
STAT_ROWS = {}


def lookup(field, term_cell, ch_cell):
    return (f'=IF({ch_cell}="","",IF(COUNTIFS({CD_TERM},{term_cell},{CD_CH},{ch_cell},{field},"<>")=0,"",'
            f'SUMIFS({field},{CD_TERM},{term_cell},{CD_CH},{ch_cell})))')


def matrix_frame(ws, title, fmt):
    put(ws, 1, 1, title, F_B)
    head(ws, 2, ["Chapter"] + peer_terms)
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 30
    for j in range(2, 2 + mcols):
        ws.column_dimensions[L(j)].width = 10.5
    for i in range(n_ch):
        r = first + i
        put(ws, r, 1, f"=IF(Lists!$D${2 + i}=\"\",\"\",Lists!$D${2 + i})", F_LINK)
    # highlight our chapter
    ws.conditional_formatting.add(
        f"A{first}:{lastc}{last}",
        FormulaRule(formula=[f"$A{first}=OurChapter"], fill=FILL_US, font=Font(name=FONT, bold=True)))
    return ws


def stat_block(ws, key, rows, fmt_map):
    """rows: list of (label, formula-builder(colletter, rng) , fmt)"""
    r0 = last + 2
    STAT_ROWS[key] = {}
    for k, (label, build, fmt) in enumerate(rows):
        r = r0 + k
        put(ws, r, 1, label, F_B, fill=FILL_SUB)
        STAT_ROWS[key][label] = r
        for j in range(2, 2 + mcols):
            col = L(j)
            rng = f"{col}${first}:{col}${last}"
            put(ws, r, j, build(col, rng), F, fmt, FILL_SUB)


def our_value(col, key_ws):
    return f"INDEX({col}${first}:{col}${last},MATCH(OurChapter,$A${first}:$A${last},0))"


enough = lambda rng: f"COUNT({rng})>=MinChapters"

# GPA matrix
ws = matrix_frame(tabs["GPA Matrix"], "IFC chapter GPA by semester (term GPA as printed on the UNL report)", GPA)
for i in range(n_ch):
    for j in range(2, 2 + mcols):
        put(ws, first + i, j, lookup(CD_GPA, f"{L(j)}$2", f"$A{first + i}"), F, GPA)
ws.conditional_formatting.add(f"B{first}:{lastc}{last}",
                              ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile",
                                             mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
stat_block(ws, "gpa", [
    ("Chapters reporting a GPA", lambda c, rng: f"=COUNT({rng})", INT),
    ("IFC average (member weighted)", lambda c, rng:
        f"=IF({enough(rng)},SUMPRODUCT({rng},'Members Matrix'!{rng})/"
        f"SUMPRODUCT(--ISNUMBER({rng}),'Members Matrix'!{rng}),\"\")", GPA),
    ("IFC median chapter", lambda c, rng: f"=IF({enough(rng)},MEDIAN({rng}),\"\")", GPA),
    ("IFC top quartile cutoff", lambda c, rng: f"=IF({enough(rng)},QUARTILE({rng},3),\"\")", GPA),
    ("Our chapter", lambda c, rng: f"=IFERROR(IF({our_value(c, ws)}=\"\",\"\",{our_value(c, ws)}),\"\")", GPA),
    ("Our IFC rank (1 = best)", lambda c, rng:
        f"=IF(OR({c}{last + 6}=\"\",NOT({enough(rng)})),\"\",COUNTIF({rng},\">\"&{c}{last + 6})+1)", INT),
    ("Our percentile (100% = best)", lambda c, rng:
        f"=IF({c}{last + 7}=\"\",\"\",IF({c}{last + 2}<=1,\"\",({c}{last + 2}-{c}{last + 7})/({c}{last + 2}-1)))", PCT),
    ("Our gap vs IFC average", lambda c, rng:
        f"=IF(OR({c}{last + 6}=\"\",{c}{last + 3}=\"\"),\"\",{c}{last + 6}-{c}{last + 3})", SIGNED),
], None)

# Members matrix
ws = matrix_frame(tabs["Members Matrix"], "IFC chapter size by semester (total members incl. new members)", INT)
for i in range(n_ch):
    for j in range(2, 2 + mcols):
        put(ws, first + i, j, lookup(CD_MEM, f"{L(j)}$2", f"$A{first + i}"), F, INT)
ws.conditional_formatting.add(f"B{first}:{lastc}{last}",
                              ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="5B9BD5"))
stat_block(ws, "mem", [
    ("Chapters reporting", lambda c, rng: f"=COUNT({rng})", INT),
    ("Total IFC men", lambda c, rng: f"=IF({enough(rng)},SUM({rng}),\"\")", INT),
    ("IFC average chapter", lambda c, rng: f"=IF({enough(rng)},AVERAGE({rng}),\"\")", "0.0"),
    ("IFC median chapter", lambda c, rng: f"=IF({enough(rng)},MEDIAN({rng}),\"\")", "0.0"),
    ("Our chapter", lambda c, rng: f"=IFERROR(IF({our_value(c, ws)}=\"\",\"\",{our_value(c, ws)}),\"\")", INT),
    ("Our size rank (1 = biggest)", lambda c, rng:
        f"=IF(OR({c}{last + 6}=\"\",NOT({enough(rng)})),\"\",COUNTIF({rng},\">\"&{c}{last + 6})+1)", INT),
    ("Our size vs median (100 = median)", lambda c, rng:
        f"=IF(OR({c}{last + 6}=\"\",{c}{last + 5}=\"\"),\"\",100*{c}{last + 6}/{c}{last + 5})", IDX),
], None)

# New members matrix
ws = matrix_frame(tabs["New Members Matrix"], "IFC new members by semester (as printed on the UNL report)", INT)
for i in range(n_ch):
    for j in range(2, 2 + mcols):
        put(ws, first + i, j, lookup(CD_NM, f"{L(j)}$2", f"$A{first + i}"), F, INT)
ws.conditional_formatting.add(f"B{first}:{lastc}{last}",
                              ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="70AD47"))
stat_block(ws, "nm", [
    ("Chapters reporting", lambda c, rng: f"=COUNT({rng})", INT),
    ("Total IFC new members", lambda c, rng: f"=IF({enough(rng)},SUM({rng}),\"\")", INT),
    ("IFC average chapter", lambda c, rng: f"=IF({enough(rng)},AVERAGE({rng}),\"\")", "0.0"),
    ("IFC median chapter", lambda c, rng: f"=IF({enough(rng)},MEDIAN({rng}),\"\")", "0.0"),
    ("Our chapter", lambda c, rng: f"=IFERROR(IF({our_value(c, ws)}=\"\",\"\",{our_value(c, ws)}),\"\")", INT),
    ("Our rank (1 = most)", lambda c, rng:
        f"=IF(OR({c}{last + 6}=\"\",NOT({enough(rng)})),\"\",COUNTIF({rng},\">\"&{c}{last + 6})+1)", INT),
    ("Our share of all IFC new members", lambda c, rng:
        f"=IF(OR({c}{last + 6}=\"\",{c}{last + 3}=\"\",{c}{last + 3}=0),\"\",{c}{last + 6}/{c}{last + 3})", "0.0%"),
], None)

def mrow(sheet, key, label):
    return STAT_ROWS[key][label]


def pull(sheet, key, label, term_cell):
    r = mrow(sheet, key, label)
    return (f"IFERROR(INDEX('{sheet}'!$B${r}:${lastc}${r},MATCH({term_cell},'{sheet}'!$B$2:${lastc}$2,0)),\"\")")


def blankify(expr):
    return f'=IF({expr}="","",{expr})'


# ---------------------------------------------------------------- TX vs IFC
ws = tabs["TX vs IFC"]
put(ws, 1, 1, "=OurChapter&\" vs. the IFC, semester by semester\"", F_T)
put(ws, 2, 1, "Everything on this tab is a formula. Change 'Our chapter' on the Lists tab to see any other house.",
    F_NOTE)
groups = [("GPA", 2, 9), ("Size", 11, 14), ("Recruitment", 15, 19), ("Retention (reported Fall 2024+)", 20, 22)]
for label, a, b in groups:
    ws.merge_cells(start_row=3, start_column=a, end_row=3, end_column=b)
    c = put(ws, 3, a, label, F_H, fill=FILL_H)
    c.alignment = CENTER
tx_cols = ["Semester",
           "Our GPA", "IFC avg (member wtd)", "IFC median", "All UNL men", "Greek men", "Gap vs IFC avg",
           "IFC GPA rank", "Chapters w/ GPA",
           "", "Our members", "IFC median size", "Size rank", "Size vs median (100 = median)",
           "Our new members", "IFC avg new members", "IFC median new members", "NM rank (1 = most)",
           "Share of all IFC NMs",
           "Our active retention", "Our NM retention", "IFC median NM retention"]
head(ws, 4, tx_cols, height=58)
ws.freeze_panes = "B5"
for i, t in enumerate(peer_terms):
    r = 5 + i
    a = f"$A{r}"
    put(ws, r, 1, t, F_LINK)
    put(ws, r, 2, blankify(pull("GPA Matrix", "gpa", "Our chapter", a)), F, GPA)
    put(ws, r, 3, blankify(pull("GPA Matrix", "gpa", "IFC average (member weighted)", a)), F, GPA)
    put(ws, r, 4, blankify(pull("GPA Matrix", "gpa", "IFC median chapter", a)), F, GPA)
    put(ws, r, 5, "=" + campus(a, "All Male"), F, GPA)
    put(ws, r, 6, "=" + campus(a, "Greek Male"), F, GPA)
    put(ws, r, 7, f'=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})', F, SIGNED)
    put(ws, r, 8, blankify(pull("GPA Matrix", "gpa", "Our IFC rank (1 = best)", a)), F, INT)
    put(ws, r, 9, f'=IF(H{r}="","",{pull("GPA Matrix", "gpa", "Chapters reporting a GPA", a)})', F, INT)
    put(ws, r, 11, blankify(pull("Members Matrix", "mem", "Our chapter", a)), F, INT)
    put(ws, r, 12, blankify(pull("Members Matrix", "mem", "IFC median chapter", a)), F, "0.0")
    put(ws, r, 13, blankify(pull("Members Matrix", "mem", "Our size rank (1 = biggest)", a)), F, INT)
    put(ws, r, 14, blankify(pull("Members Matrix", "mem", "Our size vs median (100 = median)", a)), F, IDX)
    put(ws, r, 15, blankify(pull("New Members Matrix", "nm", "Our chapter", a)), F, INT)
    put(ws, r, 16, blankify(pull("New Members Matrix", "nm", "IFC average chapter", a)), F, "0.0")
    put(ws, r, 17, blankify(pull("New Members Matrix", "nm", "IFC median chapter", a)), F, "0.0")
    put(ws, r, 18, blankify(pull("New Members Matrix", "nm", "Our rank (1 = most)", a)), F, INT)
    put(ws, r, 19, blankify(pull("New Members Matrix", "nm", "Our share of all IFC new members", a)), F, "0.0%")
    put(ws, r, 20, f'=IFERROR(IF(COUNTIFS({CD_TERM},{a},{CD_CH},OurChapter,{CD_RA},"<>")=0,"",'
                   f'SUMIFS({CD_RA},{CD_TERM},{a},{CD_CH},OurChapter)),"")', F, PCT)
    put(ws, r, 21, f'=IFERROR(IF(COUNTIFS({CD_TERM},{a},{CD_CH},OurChapter,{CD_RN},"<>")=0,"",'
                   f'SUMIFS({CD_RN},{CD_TERM},{a},{CD_CH},OurChapter)),"")', F, PCT)
    bt, bc, br = cd("A"), cd("C"), cd("K")
    ws.cell(r, 22).value = ArrayFormula(
        f"V{r}", f'=IF(COUNTIFS({CD_TERM},{a},{CD_COUNCIL},"IFC",{CD_RN},">=0")<MinChapters,"",'
                 f'MEDIAN(IF(({bt}={a})*({bc}="IFC")*ISNUMBER({br}),{br})))')
    ws.cell(r, 22).font, ws.cell(r, 22).number_format = F, PCT
TX_LAST = 4 + len(peer_terms)
for col, w in zip(range(1, 23), [12, 8, 9, 8, 8, 8, 8, 7, 8, 2, 8, 8, 7, 9, 8, 9, 9, 8, 9, 9, 9, 9]):
    ws.column_dimensions[L(col)].width = w
ws.conditional_formatting.add(f"G5:G{TX_LAST}", CellIsRule(operator="lessThan", formula=["0"],
                                                          font=Font(name=FONT, color="C00000")))
ws.conditional_formatting.add(f"G5:G{TX_LAST}", CellIsRule(operator="greaterThan", formula=["0"],
                                                          font=Font(name=FONT, color="008000")))
put(ws, TX_LAST + 1, 1, "Spring new-member counts are tiny for every house (the IFC median is 1 to 4), so compare "
                        "recruiting by calendar year on the Recruitment Index tab rather than spring by spring.", F_NOTE)

# charts
n_known = sum(1 for t in peer_terms if term_key(t) <= term_key(latest_term))
cats = Reference(ws, min_col=1, min_row=5, max_row=4 + n_known)
ch = BarChart()
ch.title, ch.height, ch.width = "Our GPA minus the IFC average (member weighted)", 8, 18
ch.y_axis.number_format = "+0.00;-0.00;0.00"
ch.legend = None
ch.x_axis.tickLblPos = "low"
ch.add_data(Reference(ws, min_col=7, min_row=4, max_row=4 + n_known), titles_from_data=True)
ch.set_categories(cats)
ws.add_chart(ch, f"A{TX_LAST + 3}")
ch2 = BarChart()
ch2.title, ch2.height, ch2.width = "New members: ours vs the average IFC house", 8, 18
for col in (15, 16):
    ch2.add_data(Reference(ws, min_col=col, min_row=4, max_row=4 + n_known), titles_from_data=True)
ch2.set_categories(cats)
ws.add_chart(ch2, f"L{TX_LAST + 3}")
ch3 = BarChart()
ch3.title, ch3.height, ch3.width = "Chapter size: ours vs the median IFC house", 8, 18
for col in (11, 12):
    ch3.add_data(Reference(ws, min_col=col, min_row=4, max_row=4 + n_known), titles_from_data=True)
ch3.set_categories(cats)
ws.add_chart(ch3, f"A{TX_LAST + 20}")
ch3b = BarChart()
ch3b.title, ch3b.height, ch3b.width = "GPA: ours vs the IFC average", 8, 18
ch3b.y_axis.scaling.min, ch3b.y_axis.scaling.max = 2.6, 3.6
ch3b.y_axis.number_format = "0.00"
for col in (2, 3):
    ch3b.add_data(Reference(ws, min_col=col, min_row=4, max_row=4 + n_known), titles_from_data=True)
ch3b.set_categories(cats)
ws.add_chart(ch3b, f"L{TX_LAST + 20}")
put(ws, TX_LAST + 2, 1, "Charts cover Fall 2018 to " + latest_term + ". After adding semesters, "
                        "extend each chart's data range (right-click > Select Data).", F_NOTE)

# ---------------------------------------------------------------- Recruitment Index (calendar years)
ws = tabs["Recruitment Index"]
put(ws, 1, 1, "Recruitment Index v2: our recruiting vs the other IFC houses, by calendar year", F_T)
put(ws, 2, 1, "Calendar year = Spring + Fall of the same year, which lines up with one exec team "
              "(elected in November, trained in December, running recruitment from January).", F_NOTE)
years = sorted({int(t.split()[1]) for t in peer_terms if t.startswith("Spring")})
groups = [("Scores (100 = matches the benchmark)", 4, 8), ("Recruiting volume", 9, 11),
          ("Size and replacement", 12, 16), ("IFC comparison inputs", 17, 21)]
for label, a, b in groups:
    ws.merge_cells(start_row=3, start_column=a, end_row=3, end_column=b)
    c = put(ws, 3, a, label, F_H, fill=FILL_H)
    c.alignment = CENTER
ri_cols = ["Year", "Spring", "Fall",
           "Recruiting index", "Size-adjusted recruiting index", "Replacement index", "Growth vs IFC (pts)",
           "Combined recruitment index",
           "Our NMs (Spring + Fall)", "IFC avg NMs per house", "Our share of all IFC NMs",
           "Our members last Fall", "Our members this Fall", "Our est. departures", "Our growth",
           "Our NMs per member",
           "IFC NMs (Spring + Fall)", "IFC men last Fall", "IFC men this Fall", "IFC growth",
           "IFC NMs per member"]
head(ws, 4, ri_cols, height=58)
ws.freeze_panes = "D5"
for i, y in enumerate(years):
    r = 5 + i
    sp, fa, pf = f"$B{r}", f"$C{r}", f'"Fall "&($A{r}-1)'
    put(ws, r, 1, y, F_IN)
    put(ws, r, 2, f'="Spring "&A{r}')
    put(ws, r, 3, f'="Fall "&A{r}')
    our_s, our_f = pull("New Members Matrix", "nm", "Our chapter", sp), pull("New Members Matrix", "nm", "Our chapter", fa)
    put(ws, r, 9, f'=IF(OR({our_s}="",{our_f}=""),"",{our_s}+{our_f})', F, INT)
    avg_s, avg_f = (pull("New Members Matrix", "nm", "IFC average chapter", x) for x in (sp, fa))
    put(ws, r, 10, f'=IF(OR({avg_s}="",{avg_f}=""),"",{avg_s}+{avg_f})', F, "0.0")
    tot_s, tot_f = (pull("New Members Matrix", "nm", "Total IFC new members", x) for x in (sp, fa))
    put(ws, r, 17, f'=IF(OR({tot_s}="",{tot_f}=""),"",{tot_s}+{tot_f})', F, INT)
    put(ws, r, 11, f'=IF(OR(I{r}="",Q{r}="",N(Q{r})=0),"",I{r}/Q{r})', F, "0.0%")
    put(ws, r, 12, blankify(pull("Members Matrix", "mem", "Our chapter", pf)), F, INT)
    put(ws, r, 13, blankify(pull("Members Matrix", "mem", "Our chapter", fa)), F, INT)
    put(ws, r, 14, f'=IF(OR(I{r}="",L{r}="",M{r}=""),"",L{r}+I{r}-M{r})', F, INT)
    put(ws, r, 15, f'=IF(OR(L{r}="",M{r}="",N(L{r})=0),"",M{r}/L{r}-1)', F, SIGNED_PCT)
    put(ws, r, 16, f'=IF(OR(I{r}="",L{r}="",N(L{r})=0),"",I{r}/L{r})', F, "0.00")
    put(ws, r, 18, blankify(pull("Members Matrix", "mem", "Total IFC men", pf)), F, INT)
    put(ws, r, 19, blankify(pull("Members Matrix", "mem", "Total IFC men", fa)), F, INT)
    put(ws, r, 20, f'=IF(OR(R{r}="",S{r}="",N(R{r})=0),"",S{r}/R{r}-1)', F, SIGNED_PCT)
    put(ws, r, 21, f'=IF(OR(Q{r}="",R{r}="",N(R{r})=0),"",Q{r}/R{r})', F, "0.00")
    put(ws, r, 4, f'=IF(OR(I{r}="",J{r}="",N(J{r})=0),"",100*I{r}/J{r})', F, IDX)
    put(ws, r, 5, f'=IF(OR(P{r}="",U{r}="",N(U{r})=0),"",100*P{r}/U{r})', F, IDX)
    put(ws, r, 6, f'=IF(OR(N{r}="",N(N{r})<=0),"",100*I{r}/N{r})', F, IDX)
    put(ws, r, 7, f'=IF(OR(O{r}="",T{r}=""),"",100*(O{r}-T{r}))', F, "+0.0;-0.0;0.0")
    put(ws, r, 8, f'=IF(COUNT(D{r}:F{r})=0,"",AVERAGE(D{r}:F{r}))', F, IDX)
RI_LAST = 4 + len(years)
for col in "DEFH":
    ws.conditional_formatting.add(f"{col}5:{col}{RI_LAST}",
                                  ColorScaleRule(start_type="num", start_value=50, start_color="F8696B",
                                                 mid_type="num", mid_value=100, mid_color="FFFFFF",
                                                 end_type="num", end_value=150, end_color="63BE7B"))
ws.conditional_formatting.add(f"G5:G{RI_LAST}", CellIsRule(operator="lessThan", formula=["0"],
                                                          font=Font(name=FONT, color="C00000")))
ws.column_dimensions["A"].width = 7
for j in range(2, 22):
    ws.column_dimensions[L(j)].width = 11 if j > 3 else 12
ri_notes = [
    "How to read it: each index is 100 when we match the benchmark. 150 is 50% better, 50 is half as good.",
    "Recruiting index = our new members for the year / the average IFC house's new members for the year. "
    "Every house recruits in the same seasons, so no fall/spring weighting is needed.",
    "Size-adjusted recruiting index = our new members per member (members = last Fall's roster) / the same "
    "ratio for the whole IFC. 100 = we recruit as hard as the IFC does for our size.",
    "Replacement index = new members / estimated departures x 100, where departures = last Fall's roster + "
    "this year's new members - this Fall's roster. 100 = we replaced every man who graduated or left.",
    "Growth vs IFC = our Fall-to-Fall % change in members minus the IFC's, in percentage points.",
    "Combined recruitment index = average of the three indexes that exist that year.",
    "Blank years: UNL left spring new members off the Spring 2020 and Spring 2022 reports, and there is no Fall 2023 "
    "report in the repo, so 2020, 2022, 2023 and 2024 are missing some IFC inputs. The current year fills in once "
    "its Fall report is added.",
    "Why v2: the old index (AE List.xlsx) only compared us to our own history, so 100 meant 'normal for Theta Xi', "
    "not 'normal for UNL'. See the Old Index Review tab.",
]
for k, text in enumerate(ri_notes):
    put(ws, RI_LAST + 2 + k, 1, text, F_NOTE)
ch6 = BarChart()
ch6.title, ch6.height, ch6.width = "Recruitment indexes by year (100 = benchmark)", 8, 20
last_full = 4 + sum(1 for y in years if f"Fall {y}" in known_terms)
for col in (4, 5, 6):
    ch6.add_data(Reference(ws, min_col=col, min_row=4, max_row=last_full), titles_from_data=True)
ch6.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last_full))
ws.add_chart(ch6, f"D{RI_LAST + 11}")

# ---------------------------------------------------------------- Combined Score
ws = tabs["Combined Score"]
put(ws, 1, 1, "Combined GPA + recruitment scorecard for every IFC house", F_T)
put(ws, 2, 1, "Pick a window and weights in the yellow cells. Scores are percentiles among IFC houses "
              "(100 = best in the IFC, 0 = worst), then blended.", F_NOTE)
_season, _year = latest_term.split()
put(ws, 3, 1, "Window start", F_B); put(ws, 3, 2, f"{_season} {int(_year) - 2}", F_IN, fill=FILL_INPUT)
put(ws, 4, 1, "Window end", F_B); put(ws, 4, 2, "=LatestTerm", F_LINK, fill=FILL_INPUT)
put(ws, 5, 1, "GPA weight", F_B); put(ws, 5, 2, 0.5, F_IN, PCT, FILL_INPUT)
put(ws, 6, 1, "Recruitment weight", F_B); put(ws, 6, 2, "=1-B5", F, PCT)
put(ws, 3, 3, f"=MATCH(B3,{TERM_LIST},0)", F)
put(ws, 4, 3, f"=MATCH(B4,{TERM_LIST},0)", F)
put(ws, 3, 4, "<- term numbers used in the formulas", F_NOTE)
put(ws, 7, 1, "Growth compares members at the window end with the window start, so start and end in the same season (Spring to Spring or Fall to Fall).", F_NOTE)
dv = DataValidation(type="list", formula1=TERM_LIST, allow_blank=False)
ws.add_data_validation(dv)
dv.add("B3"); dv.add("B4")
cs_cols = ["Chapter", "Avg GPA", "Avg members", "Avg NMs / semester", "NMs per member / semester",
           "Members at start", "Members at end", "Growth", "Avg NM retention",
           "GPA pctl", "NM volume pctl", "NM rate pctl", "Growth pctl", "Recruitment score",
           "Combined score", "Combined rank", "GPA rank", "Recruitment rank"]
head(ws, 9, cs_cols, height=45)
ws.freeze_panes = "B10"
R0, R1 = 10, 10 + n_ch - 1
win = f'{CD_ORDER},">="&$C$3,{CD_ORDER},"<="&$C$4'
for i in range(n_ch):
    r = R0 + i
    a = f"$A{r}"
    put(ws, r, 1, f"=IF(Lists!$D${2 + i}=\"\",\"\",Lists!$D${2 + i})", F_LINK)
    put(ws, r, 2, f'=IF($G{r}="","",IF(COUNTIFS({CD_CH},{a},{CD_GPA},">0",{win})=0,"",AVERAGEIFS({CD_GPA},{CD_CH},{a},{CD_GPA},">0",{win})))', F, GPA)
    put(ws, r, 3, f'=IF($G{r}="","",IF(COUNTIFS({CD_CH},{a},{CD_MEM},">0",{win})=0,"",AVERAGEIFS({CD_MEM},{CD_CH},{a},{CD_MEM},">0",{win})))', F, "0.0")
    put(ws, r, 4, f'=IF($G{r}="","",IF(COUNTIFS({CD_CH},{a},{CD_NM},">=0",{win})=0,"",AVERAGEIFS({CD_NM},{CD_CH},{a},{CD_NM},">=0",{win})))', F, "0.0")
    put(ws, r, 5, f'=IF(OR(C{r}="",D{r}=""),"",D{r}/C{r})', F, "0.00")
    put(ws, r, 6, f'=IF($G{r}="","",IF(COUNTIFS({CD_CH},{a},{CD_TERM},$B$3,{CD_MEM},">0")=0,"",SUMIFS({CD_MEM},{CD_CH},{a},{CD_TERM},$B$3)))', F, INT)
    put(ws, r, 7, f'=IF({a}="","",IF(COUNTIFS({CD_CH},{a},{CD_TERM},$B$4,{CD_MEM},">0")=0,"",SUMIFS({CD_MEM},{CD_CH},{a},{CD_TERM},$B$4)))', F, INT)
    put(ws, r, 8, f'=IF(OR(F{r}="",G{r}=""),"",G{r}/F{r}-1)', F, SIGNED_PCT)
    put(ws, r, 9, f'=IF($G{r}="","",IF(COUNTIFS({CD_CH},{a},{CD_RN},">=0",{win})=0,"",AVERAGEIFS({CD_RN},{CD_CH},{a},{CD_RN},">=0",{win})))', F, PCT)
    for col, src in zip("JKLM", "BDEH"):
        put(ws, r, ord(col) - 64,
            f'=IF({src}{r}="","",PERCENTRANK(${src}${R0}:${src}${R1},{src}{r}))', F, PCT)
    put(ws, r, 14, f'=IF(COUNT(K{r}:M{r})=0,"",AVERAGE(K{r}:M{r}))', F, PCT)
    put(ws, r, 15, f'=IF(OR(J{r}="",N{r}=""),"",$B$5*J{r}+$B$6*N{r})', F, PCT)
    put(ws, r, 16, f'=IF(O{r}="","",COUNTIF($O${R0}:$O${R1},">"&O{r})+1)', F, INT)
    put(ws, r, 17, f'=IF(B{r}="","",COUNTIF($B${R0}:$B${R1},">"&B{r})+1)', F, INT)
    put(ws, r, 18, f'=IF(N{r}="","",COUNTIF($N${R0}:$N${R1},">"&N{r})+1)', F, INT)
ws.conditional_formatting.add(f"A{R0}:R{R1}",
                              FormulaRule(formula=[f"$A{R0}=OurChapter"], fill=FILL_US, font=Font(name=FONT, bold=True)))
for col in "JKLMNO":
    ws.conditional_formatting.add(f"{col}{R0}:{col}{R1}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                 mid_type="num", mid_value=0.5, mid_color="FFFFFF",
                                                 end_type="num", end_value=1, end_color="63BE7B"))
ws.column_dimensions["A"].width = 24
for j in range(2, 19):
    ws.column_dimensions[L(j)].width = 10
cs_notes = [
    "GPA pctl: where the house's average GPA over the window falls among IFC houses.",
    "Recruitment score = average of three percentiles: new-member volume, new members per member "
    "(recruiting for its size) and membership growth across the window.",
    "Combined score = GPA weight x GPA pctl + recruitment weight x recruitment score. Ranks: 1 = best.",
    "Only houses with a report at the window end are scored (closed or suspended houses drop out). Houses that joined mid-window get no growth score.",
]
for k, text in enumerate(cs_notes):
    put(ws, R1 + 2 + k, 1, text, F_NOTE)

# ---------------------------------------------------------------- TX History (2010 on)
ws = tabs["TX History"]
put(ws, 1, 1, "=OurChapter&\" history since 2010 against the all-campus benchmarks\"", F_T)
put(ws, 2, 1, "Before Fall 2018 the repo only has our own numbers (from AE List.xlsx and GPA Data.xlsx), "
              "so peer comparisons start in Fall 2018. The IFC average size before 2018 comes from AE List.xlsx.",
    F_NOTE)
h_cols = ["Semester", "Our GPA", "Greek men", "All UNL men", "Gap vs Greek men", "Gap vs all men",
          "Our members", "Our new members", "IFC avg size", "Size vs IFC avg (100 = avg)"]
head(ws, 4, h_cols, height=45)
ws.freeze_panes = "B5"
for i, t in enumerate(all_terms):
    r = 5 + i
    a = f"$A{r}"
    put(ws, r, 1, t, F_LINK)
    put(ws, r, 2, f'=IF(COUNTIFS({CD_TERM},{a},{CD_CH},OurChapter,{CD_GPA},"<>")=0,"",'
                  f'SUMIFS({CD_GPA},{CD_TERM},{a},{CD_CH},OurChapter))', F, GPA)
    put(ws, r, 3, "=" + campus(a, "Greek Male"), F, GPA)
    put(ws, r, 4, "=" + campus(a, "All Male"), F, GPA)
    put(ws, r, 5, f'=IF(OR(B{r}="",C{r}=""),"",B{r}-C{r})', F, SIGNED)
    put(ws, r, 6, f'=IF(OR(B{r}="",D{r}=""),"",B{r}-D{r})', F, SIGNED)
    put(ws, r, 7, f'=IF(COUNTIFS({CD_TERM},{a},{CD_CH},OurChapter,{CD_MEM},"<>")=0,"",'
                  f'SUMIFS({CD_MEM},{CD_TERM},{a},{CD_CH},OurChapter))', F, INT)
    put(ws, r, 8, f'=IF(COUNTIFS({CD_TERM},{a},{CD_CH},OurChapter,{CD_NM},"<>")=0,"",'
                  f'SUMIFS({CD_NM},{CD_TERM},{a},{CD_CH},OurChapter))', F, INT)
    put(ws, r, 9, f"=IFERROR(IF({pull('Members Matrix', 'mem', 'IFC average chapter', a)}<>\"\","
                  f"{pull('Members Matrix', 'mem', 'IFC average chapter', a)},"
                  f"IF(INDEX('IFC Size (legacy)'!$B:$B,MATCH({a},'IFC Size (legacy)'!$A:$A,0))=\"\",\"\","
                  f"INDEX('IFC Size (legacy)'!$B:$B,MATCH({a},'IFC Size (legacy)'!$A:$A,0)))),\"\")", F, "0.0")
    put(ws, r, 10, f'=IF(OR(G{r}="",I{r}=""),"",100*G{r}/I{r})', F, IDX)
H_LAST = 4 + len(all_terms)
for col, w in zip(range(1, 11), [12, 9, 9, 9, 9, 9, 9, 9, 9, 11]):
    ws.column_dimensions[L(col)].width = w
n_hist = sum(1 for t in all_terms if term_key(t) <= term_key(latest_term))
ch4 = BarChart()
ch4.title, ch4.height, ch4.width = "Our GPA minus the Greek men GPA, since 2010", 8, 20
ch4.y_axis.number_format = "+0.00;-0.00;0.00"
ch4.legend = None
ch4.x_axis.tickLblPos = "low"
ch4.add_data(Reference(ws, min_col=5, min_row=4, max_row=4 + n_hist), titles_from_data=True)
ch4.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + n_hist))
ws.add_chart(ch4, "L4")
ch5 = BarChart()
ch5.title, ch5.height, ch5.width = "Our members and new members since 2010", 8, 20
for col in (7, 8):
    ch5.add_data(Reference(ws, min_col=col, min_row=4, max_row=4 + n_hist), titles_from_data=True)
ch5.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + n_hist))
ws.add_chart(ch5, "L21")

# legacy IFC average size (AE List.xlsx) lives on its own small tab
wsl = tabs["IFC Size (legacy)"]
head(wsl, 1, ["Semester", "IFC avg chapter size", "Source"])
for i, r in enumerate([r for r in campus_rows if num(r.get("IFC avg chapter size"))], start=2):
    put(wsl, i, 1, r["semester"], F_IN)
    put(wsl, i, 2, num(r["IFC avg chapter size"]), F_IN, "0.0")
    put(wsl, i, 3, "AE List.xlsx, Cleaned tab, column K (user's own records)", F_IN)
wsl.column_dimensions["A"].width = 13
wsl.column_dimensions["B"].width = 12
wsl.column_dimensions["C"].width = 55

# pledge-class retention (AE List.xlsx, Disgusting tab) below the history table
ws = tabs["TX History"]
pr = H_LAST + 3
put(ws, pr, 1, "Theta Xi pledge-class retention (from AE List.xlsx, Disgusting tab, top block)", F_B)
head(ws, pr + 1, ["Class", "Pledged", "Still in", "Lost", "% lost"])
classes = [("Fall 2020", 20, 10), ("Spring 2021", 10, 6), ("Fall 2021", 13, 9), ("Spring 2022", 5, 4),
           ("Fall 2022", 13, 11), ("Spring 2023", 5, 3), ("Fall 2023", 20, 13)]
for k, (cls, size, kept) in enumerate(classes):
    r = pr + 2 + k
    put(ws, r, 1, cls, F_IN)
    put(ws, r, 2, size, F_IN, INT)
    put(ws, r, 3, kept, F_IN, INT)
    put(ws, r, 4, f"=B{r}-C{r}", F, INT)
    put(ws, r, 5, f'=IF(B{r}=0,"",D{r}/B{r})', F, PCT)
rt = pr + 2 + len(classes)
put(ws, rt, 1, "Total", F_B)
for col in "BCD":
    put(ws, rt, ord(col) - 64, f"=SUM({col}{pr + 2}:{col}{rt - 1})", F_B, INT)
put(ws, rt, 5, f"=D{rt}/B{rt}", F_B, PCT)
put(ws, rt + 1, 1, "Counted when that sheet was last updated (late 2023). The chart table lower on that sheet "
                   "shows slightly different splits for Fall 2020 and Spring 2023. From Fall 2024 on, UNL reports "
                   "new-member retention for every house (see TX vs IFC).", F_NOTE)

# ---------------------------------------------------------------- Summary
ws = tabs["Summary"]
ws.column_dimensions["A"].width = 44
for col in "BCDE":
    ws.column_dimensions[col].width = 16
put(ws, 1, 1, '=OurChapter&" at UNL: GPA and recruitment compared with the other IFC houses"', F_T)
put(ws, 2, 1, "The tables are live formulas. The written evaluation below them was written from data through "
              + latest_term + ".", F_NOTE)
put(ws, 4, 1, "Latest semester", F_B); put(ws, 4, 2, "=LatestTerm", F_LINK)
put(ws, 5, 1, "Recruitment year to show (calendar year)", F_B)
put(ws, 5, 2, max(int(t.split()[1]) for t in known_terms if t.startswith("Fall")), F_IN, fill=FILL_INPUT)
put(ws, 5, 3, "Pick the last year with both a Spring and a Fall report.", F_NOTE)

head(ws, 7, ["Latest semester", "Us", "IFC benchmark", "Our IFC rank", "Houses ranked"])


def txv(col):
    return f"=IFERROR(INDEX('TX vs IFC'!${col}$5:${col}${TX_LAST},MATCH(LatestTerm,'TX vs IFC'!$A$5:$A${TX_LAST},0)),\"\")"


snap = [("Chapter GPA (benchmark = IFC average, member weighted)", "B", "C", "H", "I", GPA),
        ("Members (benchmark = IFC median house)", "K", "L", "M", "I", INT),
        ("New members this semester (benchmark = IFC average house)", "O", "P", "R", "I", "0.0"),
        ("New-member retention (benchmark = IFC median)", "U", "V", None, None, PCT)]
for k, (label, us, bench, rank, n, fmt) in enumerate(snap):
    r = 8 + k
    put(ws, r, 1, label, F)
    put(ws, r, 2, txv(us), F_LINK, fmt)
    put(ws, r, 3, txv(bench), F_LINK, fmt)
    if rank:
        put(ws, r, 4, txv(rank), F_LINK, INT)
        put(ws, r, 5, txv(n), F_LINK, INT)

head(ws, 13, ["Recruitment for the chosen calendar year", "Us", "Benchmark", "", ""])


def riv(col):
    return f"=IFERROR(INDEX('Recruitment Index'!${col}$5:${col}${RI_LAST},MATCH($B$5,'Recruitment Index'!$A$5:$A${RI_LAST},0)),\"\")"


ann = [("New members (Spring + Fall) vs the average IFC house", riv("I"), riv("J"), "0.0"),
       ("Recruiting index (100 = average IFC house)", riv("D"), 100, IDX),
       ("Size-adjusted recruiting index (100 = IFC rate)", riv("E"), 100, IDX),
       ("Replacement index (100 = replaced everyone who left)", riv("F"), 100, IDX),
       ("Membership growth, Fall to Fall (us vs the IFC)", riv("O"), riv("T"), SIGNED_PCT),
       ("Share of all IFC new members", riv("K"), f'=IFERROR(1/INDEX(\'New Members Matrix\'!$B${STAT_ROWS["nm"]["Chapters reporting"]}:${lastc}${STAT_ROWS["nm"]["Chapters reporting"]},MATCH("Fall "&$B$5,\'New Members Matrix\'!$B$2:${lastc}$2,0)),"")', "0.0%")]
for k, (label, us, bench, fmt) in enumerate(ann):
    r = 14 + k
    put(ws, r, 1, label, F)
    put(ws, r, 2, us, F_LINK, fmt)
    put(ws, r, 3, bench, F_LINK if isinstance(bench, str) else F, fmt)
put(ws, 19, 4, "<- benchmark = an equal share (1 / number of houses)", F_NOTE)

head(ws, 21, ["Combined scorecard (window on the Combined Score tab)", "Our rank", "Houses scored", "", ""])
cs_n = f"COUNT('Combined Score'!$O${R0}:$O${R1})"


def csv_(col):
    return f"=IFERROR(INDEX('Combined Score'!${col}${R0}:${col}${R1},MATCH(OurChapter,'Combined Score'!$A${R0}:$A${R1},0)),\"\")"


for k, (label, col) in enumerate([("GPA rank", "Q"), ("Recruitment rank", "R"), ("Combined rank", "P")]):
    r = 22 + k
    put(ws, r, 1, label, F)
    put(ws, r, 2, csv_(col), F_LINK, INT)
    put(ws, r, 3, "=" + cs_n, F_LINK, INT)
put(ws, 25, 1, "Window", F); put(ws, 25, 2, "='Combined Score'!B3&\" to \"&'Combined Score'!B4", F_LINK)

head(ws, 27, ["Targets for the current exec team", "Value", "", "", ""])
put(ws, 28, 1, "GPA points we are behind the IFC average (latest semester)", F)
put(ws, 28, 2, '=IF(OR(B8="",C8=""),"",C8-B8)', F, SIGNED)
put(ws, 29, 1, "New members this Fall to match last year's average IFC house", F)
this_year = int(latest_term.split()[1])
put(ws, 29, 2, f'=IFERROR(MAX(0,ROUNDUP({riv("J")[1:]}-INDEX(\'New Members Matrix\'!$B${STAT_ROWS["nm"]["Our chapter"]}:${lastc}${STAT_ROWS["nm"]["Our chapter"]},MATCH("Spring "&($B$5+1),\'New Members Matrix\'!$B$2:${lastc}$2,0)),0)),"")', F, INT)
put(ws, 29, 3, '="(the " & $B$5 & " IFC average per house, minus our Spring " & ($B$5+1) & " new members)"', F_NOTE)
put(ws, 30, 1, "New members this Fall to replace last year's departures", F)
put(ws, 30, 2, f'=IFERROR(MAX(0,{riv("N")[1:]}-INDEX(\'New Members Matrix\'!$B${STAT_ROWS["nm"]["Our chapter"]}:${lastc}${STAT_ROWS["nm"]["Our chapter"]},MATCH("Spring "&($B$5+1),\'New Members Matrix\'!$B$2:${lastc}$2,0))),"")', F, INT)
put(ws, 30, 3, "(assumes as many men leave this year as left last year)", F_NOTE)

put(ws, 32, 1, "Evaluation", F_T)
for k, text in enumerate(EVALUATION):
    c = put(ws, 33 + k, 1, text, F)
    ws.merge_cells(start_row=33 + k, start_column=1, end_row=33 + k, end_column=5)
    c.alignment = WRAP
    ws.row_dimensions[33 + k].height = 15 * max(2, len(text) // 95 + 1)

# ---------------------------------------------------------------- Old Index Review
ws = tabs["Old Index Review"]
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110
put(ws, 1, 1, "Review of the old recruitment score and GPA math", F_T)
head(ws, 3, ["Where", "What was off, and what v2 does instead"])
for k, (where, text) in enumerate(OLD_INDEX_NOTES):
    r = 4 + k
    put(ws, r, 1, where, F_B).alignment = WRAP
    c = put(ws, r, 2, text, F)
    c.alignment = WRAP
    ws.row_dimensions[r].height = 15 * max(2, len(text) // 120 + 1)

# ---------------------------------------------------------------- Read Me
ws = tabs["Read Me"]
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 110
put(ws, 1, 1, "Theta Xi vs UNL fraternities: GPA and recruitment", F_T)
r = 3
for section, rows in README:
    put(ws, r, 1, section, F_B, fill=FILL_SUB)
    ws.cell(r, 2).fill = FILL_SUB
    r += 1
    for a, text in rows:
        put(ws, r, 1, a, F_B).alignment = WRAP
        c = put(ws, r, 2, text, F)
        c.alignment = WRAP
        ws.row_dimensions[r].height = 15 * max(1, len(text) // 120 + 1)
        r += 1
    r += 1
legend = [("Blue text", F_IN, None, "typed-in data"), ("Black text", F, None, "formula"),
          ("Green text", F_LINK, None, "formula that pulls from another tab"),
          ("Yellow fill", F_IN, FILL_INPUT, "setting you can change"),
          ("Orange fill", F, FILL_GAP, "number missing from the repo; fill it in when you find it"),
          ("Light gold row", F_B, FILL_US, "our chapter (set on the Lists tab)")]
put(ws, r, 1, "Color key", F_B, fill=FILL_SUB); ws.cell(r, 2).fill = FILL_SUB
for k, (label, font, fill, text) in enumerate(legend):
    put(ws, r + 1 + k, 1, label, font, fill=fill)
    put(ws, r + 1 + k, 2, text, F)

from openpyxl.worksheet.properties import PageSetupProperties
for w in wb.worksheets:
    for chart in w._charts:          # keep axes visible in Excel
        chart.x_axis.delete = False
        chart.y_axis.delete = False
    w.sheet_view.showGridLines = w.title in ("Chapter Data", "Campus Stats", "Lists", "IFC Size (legacy)")
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
wb.active = 0

wb.save(OUT)
print("rows", len(chapter_rows), "chapters", len(ifc_chapters), "terms", len(all_terms), "->", OUT)
